#!/usr/bin/env python3
"""Render the validated N2 MEE v4 anonymous manuscript Markdown to review DOCX.

This is a submission-format transform only. It does not alter the scientific source,
rerun empirical endpoints, or change any frozen decision rule.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)


def _suppress_line_number(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:suppressLineNumbers")):
        p_pr.remove(old)
    p_pr.append(OxmlElement("w:suppressLineNumbers"))


def _set_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _suppress_line_number(paragraph)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _all_main_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        yield from _iter_table_paragraphs(table)


def _configure_section(section) -> None:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.different_first_page_header_footer = False

    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:lnNumType")):
        sect_pr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:distance"), "360")
    ln.set(qn("w:restart"), "continuous")
    sect_pr.append(ln)

    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), "1")

    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.clear()
    _set_page_field(paragraph)


def _format_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.color.rgb = BLACK
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_after = Pt(0)

    for paragraph in _all_main_paragraphs(document):
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.color.rgb = BLACK

    for section in document.sections:
        _configure_section(section)

    props = document.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.keywords = ""
    props.subject = ""


def render(markdown: Path, output: Path, manifest: Path | None = None) -> dict[str, object]:
    if not markdown.exists():
        raise FileNotFoundError(markdown)
    output.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="n2-mee-docx-") as temp_dir:
        raw_docx = Path(temp_dir) / "pandoc.docx"
        subprocess.run(
            ["pandoc", str(markdown), "--standalone", "--from=markdown", "--to=docx", "-o", str(raw_docx)],
            check=True,
        )
        document = Document(raw_docx)
        _format_document(document)
        document.save(output)

    payload = output.read_bytes()
    result = {
        "schema_version": 1,
        "role": "n2_mee_anonymous_review_docx_v4",
        "source_markdown": markdown.as_posix(),
        "output": output.as_posix(),
        "sha256": hashlib.sha256(payload).hexdigest(),
        "bytes": len(payload),
        "single_column": True,
        "double_line_spacing": True,
        "continuous_line_numbering": True,
        "page_numbering": True,
        "footer_line_number_suppressed": True,
        "review_text_color_black": True,
        "author_metadata_blank": True,
        "scientific_source_modified": False,
        "empirical_endpoint_rerun": False,
    }
    if manifest is not None:
        manifest.parent.mkdir(parents=True, exist_ok=True)
        manifest.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--manifest", type=Path)
    args = parser.parse_args()
    print(json.dumps(render(args.input, args.output, args.manifest), sort_keys=True))


if __name__ == "__main__":
    main()
